#!/usr/bin/env python3
"""Makes the writer template (templates/The-Truth_Writer-Template.docx).
Run:  python3 tools/make_template.py [output.docx]
Needs python-docx. Writers never run this; it is here so the template can be rebuilt."""
import json
import os
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "templates", "The-Truth_Writer-Template.docx")
BOOKS = [b["name"] for b in json.load(open(os.path.join(ROOT, "tools", "books.json"), encoding="utf-8"))]

INK = RGBColor(0x1B, 0x1C, 0x18)
GREY = RGBColor(0x6B, 0x6B, 0x63)
ACCENT = RGBColor(0xB5, 0x5A, 0x42)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def shade(style, fill):
    pPr = style.element.get_or_add_pPr()
    pPr.append(parse_xml('<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls("w"), fill)))


def left_border(style, color, size=18):
    pPr = style.element.get_or_add_pPr()
    pPr.append(parse_xml('<w:pBdr %s><w:left w:val="single" w:sz="%d" w:space="8" w:color="%s"/></w:pBdr>' % (nsdecls("w"), size, color)))


def pstyle(name, base="Normal"):
    s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = styles[base]
    s.quick_style = True
    return s


for name, size in (("Heading 2", 17), ("Heading 3", 14)):
    h = styles[name]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = INK
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
t = styles["Title"]
t.font.name = "Calibri"
t.font.size = Pt(26)
t.font.bold = True
t.font.color.rgb = INK
pPr = t.element.get_or_add_pPr()
for b in pPr.findall(qn("w:pBdr")):
    pPr.remove(b)

note = pstyle("Writer Note")
note.font.size = Pt(11)
note.font.italic = True
note.font.color.rgb = RGBColor(0x5A, 0x4A, 0x10)
shade(note, "FFF4CC")
note.paragraph_format.space_after = Pt(4)
note.paragraph_format.left_indent = Inches(0.1)
note.paragraph_format.right_indent = Inches(0.1)

ov = pstyle("Overview")
ov.font.size = Pt(12)
shade(ov, "EAF0F6")
left_border(ov, "4A7BA6")
ov.paragraph_format.left_indent = Inches(0.2)

sc = pstyle("Scripture")
sc.font.italic = True
sc.font.size = Pt(13)
left_border(sc, "B55A42")
sc.paragraph_format.left_indent = Inches(0.3)
sc.paragraph_format.space_after = Pt(2)

sr = pstyle("Scripture Reference")
sr.font.size = Pt(10.5)
sr.font.bold = True
sr.font.color.rgb = ACCENT
left_border(sr, "B55A42")
sr.paragraph_format.left_indent = Inches(0.3)
sr.paragraph_format.space_after = Pt(10)

cn = pstyle("Closing Note")
cn.font.bold = True
shade(cn, "F3EDE3")
left_border(cn, "B55A42", 24)
cn.paragraph_format.left_indent = Inches(0.2)


def para(text, style="Normal", bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


# ----------------------------------------------------------------- page top
doc.add_paragraph("The Truth: Writer Template", style="Title")
para("Read this box, then delete nothing. It is not published.", "Writer Note", "How to use this template. ")
para("1. Save a copy of this file first (File > Save a Copy). Name it the book and chapter, such as John 3.docx, "
     "or the article title, such as Why Study the Bible.docx.", "Writer Note")
para("2. Fill in the grey boxes in the table below. Click a box and choose from the list, pick a date, or type.", "Writer Note")
para("3. Write your text under the line that says START WRITING BELOW. Replace each line that begins with SAMPLE, "
     "or delete it. Choose a style from the Styles gallery on the Home tab: Heading 2 for a main subheading, "
     "Heading 3 for a verse group such as Verses 1-4, Scripture for a quoted scripture followed by Scripture Reference "
     "for its verse, Overview for the chapter summary box, Closing Note for a closing line. Everything else is normal text.", "Writer Note")
para("4. Save the file and upload it to the Truth website's submissions folder. The page appears on the site in about five minutes. "
     "To change a page later, edit your file and upload it again with the same file name.", "Writer Note")
para("To keep a page unpublished while you work, set Status to Keep as draft.", "Writer Note")
doc.add_paragraph()

# ------------------------------------------------------------------- table
SDT = "w:sdt"


def sdt_xml(alias, tag, kind, prompt, items=None):
    placeholder_run = ('<w:r><w:rPr><w:color w:val="7A7A72"/></w:rPr><w:t>%s</w:t></w:r>' % prompt)
    if kind == "dropdown":
        control = "<w:dropDownList>%s</w:dropDownList>" % "".join(
            '<w:listItem w:displayText="%s" w:value="%s"/>' % (i.replace("&", "&amp;"), i.replace("&", "&amp;")) for i in items)
    elif kind == "date":
        control = ('<w:date><w:dateFormat w:val="MMMM d, yyyy"/><w:lid w:val="en-US"/>'
                   '<w:storeMappedDataAs w:val="dateTime"/><w:calendar w:val="gregorian"/></w:date>')
    else:
        control = "<w:text/>"
    return parse_xml(
        '<w:sdt %s><w:sdtPr><w:alias w:val="%s"/><w:tag w:val="%s"/><w:showingPlcHdr/>%s</w:sdtPr>'
        '<w:sdtContent>%s</w:sdtContent></w:sdt>' % (nsdecls("w"), alias, tag, control, placeholder_run))


ROWS = [
    ("Where will this be published?", "Choose where this goes", "where", "dropdown",
     ["Know the Truth (article)", "Seek Jehovah (article)", "The Prophetic Word (Bible chapter)"],
     "Choose one. This decides where the page goes on the site."),
    ("Bible book", "Choose a book (chapters only)", "book", "dropdown", BOOKS,
     "Bible chapters only. Leave empty for articles."),
    ("Chapter number", "Type the number (chapters only)", "chapter", "text", None,
     "Bible chapters only, such as 3. The title is made for you."),
    ("Title", "Type the article title (articles only)", "title", "text", None,
     "Articles only. Leave empty for Bible chapters."),
    ("Date first posted", "Click to pick a date", "posted", "date", None,
     "The day this page first goes live."),
    ("Date last updated", "Click to pick a date", "updated", "date", None,
     "The same day for a new page. Change it when you edit the page later."),
    ("Summary", "Type one sentence (articles)", "summary", "text", None,
     "Articles only. One or two sentences that appear in the list."),
    ("Status", "Choose", "status", "dropdown", ["Ready to publish", "Keep as draft"],
     "Ready to publish puts it on the site. Keep as draft holds it back."),
]

table = doc.add_table(rows=len(ROWS), cols=2)
table.style = "Table Grid"
table.autofit = False
widths = (Inches(2.5), Inches(4.0))
for i, (label, prompt, tag, kind, items, hint) in enumerate(ROWS):
    row = table.rows[i]
    for c, w in zip(row.cells, widths):
        c.width = w
    lab, val = row.cells
    lab._tc.get_or_add_tcPr().append(parse_xml('<w:shd %s w:val="clear" w:color="auto" w:fill="EDEDE8"/>' % nsdecls("w")))
    p = lab.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(11.5)
    p2 = lab.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(hint)
    r2.font.size = Pt(9)
    r2.font.color.rgb = GREY
    vp = val.paragraphs[0]
    vp.paragraph_format.space_after = Pt(0)
    vp._p.append(sdt_xml(label, tag, kind, prompt, items))

doc.add_paragraph()
line = doc.add_paragraph(style="Writer Note")
line.add_run("START WRITING BELOW THIS LINE").bold = True

# -------------------------------------------------------------------- body
para("SAMPLE: Overview. Write two or three sentences that sum up the chapter. (Bible chapters only. Delete this line if you do not want a summary box.)", "Overview")
doc.add_paragraph("SAMPLE: Main subheading (Heading 2)", style="Heading 2")
doc.add_paragraph("SAMPLE: Verses 1-4 (Heading 3, for Bible chapters)", style="Heading 3")
para("SAMPLE: “Type the scripture here, in quotation marks.”", "Scripture")
para("SAMPLE: Verse 1", "Scripture Reference")
para("SAMPLE: Type your comments on those verses here, as normal paragraphs. Bold, italic and links all carry over.", "Normal")
para("SAMPLE: An optional closing line for the reader.", "Closing Note")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved", OUT)
